"""
Document ingestion: pulls raw text out of PDF or DOCX contract files.
Scanned/image PDFs are NOT OCR'd here — add an OCR fallback (e.g. pytesseract)
if you expect scanned documents.
"""

import os
import fitz  # PyMuPDF
from docx import Document


def extract_text(file_path: str) -> str:
    """Extract raw text from a .pdf or .docx contract file."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return _extract_pdf(file_path)
    elif ext == ".docx":
        return _extract_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}. Use .pdf or .docx")


def _extract_pdf(file_path: str) -> str:
    text_parts = []
    with fitz.open(file_path) as doc:
        for page_num, page in enumerate(doc, start=1):
            page_text = page.get_text()
            if page_text.strip():
                text_parts.append(f"\n--- Page {page_num} ---\n{page_text}")
            else:
                text_parts.append(f"\n--- Page {page_num} [NO EXTRACTABLE TEXT — likely scanned, needs OCR] ---\n")
    return "\n".join(text_parts)


def _extract_docx(file_path: str) -> str:
    doc = Document(file_path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)
    return "\n".join(parts)


def chunk_text(text: str, max_chars: int = 60000) -> list:
    """
    Split long contracts into chunks that fit comfortably in one Gemini call.
    """
    if len(text) <= max_chars:
        return [text]

    chunks = []
    current = ""
    for block in text.split("\n\n"):
        if len(current) + len(block) > max_chars and current:
            chunks.append(current)
            current = block
        else:
            current += "\n\n" + block if current else block
    if current:
        chunks.append(current)
    return chunks