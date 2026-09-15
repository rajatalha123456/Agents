"""
§6.1 — Excel (xlsx) parser. §17.1: "System header detect karta hai" and
"Multi-sheet, merged header handling" are both this module's job; column
*meaning* is still a mapping-studio concern, same as CSV — this parser has
no idea which column is the amount, only where the header row and the data
start.

Deliberately produces the exact same `ParsedFile`/`RawRow` shape as
`csv_parser` (string-only field values) so the rest of the ingestion
pipeline never has to branch on source format.
"""
from __future__ import annotations

import io
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from core.ingestion.parsers.base import ParsedFile, ParseError, RawRow

_LOAD_ERRORS = (InvalidFileException, KeyError, OSError, zipfile.BadZipFile)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):  # date/datetime
        return value.isoformat()
    return str(value)


def list_sheet_names(raw_bytes: bytes) -> list[str]:
    """§17.1 multi-sheet support: lets the mapping studio offer a sheet
    picker before committing to one, instead of silently guessing.
    """
    try:
        workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except _LOAD_ERRORS as exc:
        raise ParseError(f"could not read Excel file: {exc}") from exc
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _find_header_row(rows: list[tuple]) -> int:
    for idx, row in enumerate(rows):
        if any(cell is not None and str(cell).strip() for cell in row):
            return idx
    raise ParseError("no header row detected")


def _resolve_merged_headers(header_cells: tuple) -> list[str]:
    """A merged header cell (e.g. a title spanning 3 columns) stores its
    text only in the leftmost cell of the span; openpyxl reports every
    other cell in that span as None. Forward-fill from the last non-blank
    header so every column still gets a real name, then disambiguate any
    resulting duplicates (common for merged spans) with a numeric suffix —
    never leave a column with an empty or colliding header.
    """
    filled: list[str] = []
    last = ""
    for cell in header_cells:
        text = str(cell).strip() if cell is not None and str(cell).strip() else ""
        if text:
            last = text
        filled.append(last or "COLUMN")

    seen: dict[str, int] = {}
    result: list[str] = []
    for name in filled:
        seen[name] = seen.get(name, 0) + 1
        result.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    return result


def parse_excel(raw_bytes: bytes, sheet_name: str | None = None) -> ParsedFile:
    try:
        workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except _LOAD_ERRORS as exc:
        raise ParseError(f"could not read Excel file: {exc}") from exc

    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise ParseError(f"sheet {sheet_name!r} not found in workbook")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook[workbook.sheetnames[0]]

        all_rows = list(sheet.iter_rows(values_only=True))
        if not all_rows:
            raise ParseError("sheet is empty")

        header_idx = _find_header_row(all_rows)
        headers = _resolve_merged_headers(all_rows[header_idx])

        rows: list[RawRow] = []
        for offset, raw_row in enumerate(all_rows[header_idx + 1:]):
            if all(cell is None or str(cell).strip() == "" for cell in raw_row):
                continue  # trailing/interstitial blank row, not data
            values = list(raw_row) + [None] * (len(headers) - len(raw_row))
            fields = {header: _cell_text(value) for header, value in zip(headers, values)}
            # openpyxl rows are 1-indexed; header_idx is 0-indexed, so the
            # first data row's 1-indexed sheet row is header_idx + 2.
            rows.append(RawRow(line_no=header_idx + 2 + offset, fields=fields))

        return ParsedFile(rows=rows, source_format="EXCEL")
    finally:
        workbook.close()
