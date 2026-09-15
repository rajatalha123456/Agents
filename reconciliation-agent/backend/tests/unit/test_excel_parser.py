from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from core.ingestion.parsers.base import ParseError
from core.ingestion.parsers.excel_parser import list_sheet_names, parse_excel


def _workbook_bytes(build) -> bytes:
    wb = Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parses_first_sheet_by_default():
    def build(wb):
        ws = wb.active
        ws.append(["Date", "Amount", "Reference"])
        ws.append(["2026-09-01", 100.5, "INV-1"])
        ws.append(["2026-09-02", -50, "INV-2"])

    result = parse_excel(_workbook_bytes(build))
    assert result.source_format == "EXCEL"
    assert len(result.rows) == 2
    assert result.rows[0].fields["Reference"] == "INV-1"
    assert result.rows[0].fields["Amount"] == "100.5"
    assert result.rows[0].line_no == 2


def test_skips_blank_leading_rows_to_find_header():
    def build(wb):
        ws = wb.active
        ws.append([None, None])
        ws.append(["Date", "Amount"])
        ws.append(["2026-09-01", 10])

    result = parse_excel(_workbook_bytes(build))
    assert len(result.rows) == 1
    assert result.rows[0].fields == {"Date": "2026-09-01", "Amount": "10"}


def test_skips_trailing_blank_rows():
    def build(wb):
        ws = wb.active
        ws.append(["Date", "Amount"])
        ws.append(["2026-09-01", 10])
        ws.append([None, None])

    result = parse_excel(_workbook_bytes(build))
    assert len(result.rows) == 1


def test_merged_header_forward_fills_and_disambiguates():
    def build(wb):
        ws = wb.active
        ws.append(["Statement", None, "Amount"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        ws.append(["2026-09-01", "INV-1", 10])

    result = parse_excel(_workbook_bytes(build))
    headers = set(result.rows[0].fields)
    assert headers == {"Statement", "Statement (2)", "Amount"}


def test_multi_sheet_picker():
    def build(wb):
        wb.active.title = "Bank"
        wb.active.append(["Date", "Amount"])
        ledger = wb.create_sheet("Ledger")
        ledger.append(["Date", "Amount"])
        ledger.append(["2026-09-01", 5])

    raw = _workbook_bytes(build)
    assert list_sheet_names(raw) == ["Bank", "Ledger"]
    result = parse_excel(raw, sheet_name="Ledger")
    assert len(result.rows) == 1
    assert result.rows[0].fields["Amount"] == "5"


def test_unknown_sheet_name_raises():
    def build(wb):
        wb.active.append(["Date"])

    with pytest.raises(ParseError):
        parse_excel(_workbook_bytes(build), sheet_name="DoesNotExist")


def test_empty_sheet_raises():
    def build(wb):
        pass

    with pytest.raises(ParseError):
        parse_excel(_workbook_bytes(build))


def test_not_an_excel_file_raises():
    with pytest.raises(ParseError):
        parse_excel(b"not an excel file at all")
